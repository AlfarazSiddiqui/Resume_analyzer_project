"""
Text extraction utilities for uploaded resume files (.pdf, .docx, .txt).
"""
import os


def extract_text(file_path: str) -> str:
    """Detect file type by extension and extract raw text."""
    ext = os.path.splitext(file_path)[1].lower()

    if ext == '.pdf':
        return _extract_pdf(file_path)
    elif ext == '.docx':
        return _extract_docx(file_path)
    elif ext == '.txt':
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    else:
        raise ValueError(f"Unsupported file type: {ext}. Please upload a PDF, DOCX, or TXT file.")


def _extract_pdf(file_path: str) -> str:
    text_parts = []
    try:
        import pdfplumber
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
    except Exception:
        # Fallback to pypdf if pdfplumber fails on a malformed PDF
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def _extract_docx(file_path: str) -> str:
    import docx
    doc = docx.Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also pull text from any tables in the résumé
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)
    return "\n".join(paragraphs)
